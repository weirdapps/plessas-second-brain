"""Text inside grouped shapes reaches the PPTX extraction.

A python-pptx GroupShape has neither a text frame nor a table, so walking only
slide.shapes dropped every label, callout and diagram box built as a group.
Sampled decks on the producer lost grouped text in a quarter to a half of the
files, while the row still said 'extracted'. Finding attachments-6.
"""

from pptx import Presentation
from pptx.util import Inches

from src.extract.attachment_extractors import extract_text_from_file

_FILLER = "Top level textbox with enough ordinary words to clear the noise filter."


def _deck(path) -> str:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.5), Inches(0.5), Inches(6), Inches(1)
    ).text_frame.text = _FILLER

    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(2), Inches(3), Inches(1)
    ).text_frame.text = "GROUPEDCALLOUT"
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(
        Inches(1), Inches(3), Inches(3), Inches(1)
    ).text_frame.text = "NESTEDGROUPLABEL"
    # GroupShapes has no add_table, so build the table on the slide and move
    # its graphicFrame into the nested group, which is how PowerPoint stores it.
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(1))
    frame.table.cell(0, 0).text = "GROUPEDTABLECELL"
    frame.table.cell(0, 1).text = "second"
    nested._element.append(frame._element)

    out = str(path / "grouped.pptx")
    prs.save(out)
    return out


def test_grouped_and_nested_group_text_is_extracted(tmp_path):
    result = extract_text_from_file(
        _deck(tmp_path),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert result["status"] == "extracted"
    assert "Top level textbox" in result["text"]
    assert "GROUPEDCALLOUT" in result["text"]
    assert "NESTEDGROUPLABEL" in result["text"]
    assert "GROUPEDTABLECELL | second" in result["text"]
