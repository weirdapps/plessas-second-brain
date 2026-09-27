"""An attachment whose name does not say what it is is read by its bytes.

Outlook saves some parts with no extension at all ("Outlook-abc12"), and users
lose the dot in a name ("Status Updatepptx"). mimetypes.guess_type returns None
for both, so the registrar recorded application/octet-stream and the extractor
skipped them as 'Unsupported type': board minutes, PDFs and screenshots sat on
disk and never reached search or the image pipeline. Older rows carry the
opposite mislabel, a real .docx declared application/zip, which the skip list
caught before the extension had a say. Finding attachments-4.
"""

import io
import struct
import zipfile
from pathlib import Path

import pytest

from src.extract import attachment_extractors as ax
from src.extract.attachment_extractors import extract_text_from_file, sniff_mime_type

_BODY = "Synthetic quarterly review with enough words to clear the noise filter easily."

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docx_bytes() -> bytes:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph(_BODY)
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    buf = io.BytesIO()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = _BODY
    prs.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active["A1"] = _BODY
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes:
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), _BODY)
    return doc.tobytes()


def _plain_zip_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("notes/readme.txt", _BODY)
    return buf.getvalue()


def _ole_bytes(stream_name: str) -> bytes:
    """A minimal OLE2 header whose first directory sector names one stream.

    Enough for a sniffer that reads the header fields the spec fixes (sector
    shift at offset 30, first directory sector at offset 48, a 128-byte entry
    with its name length at offset 64), which is all the registrar needs to tell a Word document from a mail item or a workbook.
    """
    header = bytearray(512)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<H", header, 30, 9)  # 512-byte sectors
    struct.pack_into("<I", header, 48, 0)  # directory starts at sector 0
    directory = bytearray(512)
    for offset, entry in ((0, "Root Entry"), (128, stream_name)):
        name = (entry + "\0").encode("utf-16-le")
        directory[offset : offset + len(name)] = name
        struct.pack_into("<H", directory, offset + 64, len(name))
    return bytes(header) + bytes(directory)


def _write(tmp_path: Path, name: str, data: bytes) -> str:
    path = tmp_path / name
    path.write_bytes(data)
    return str(path)


@pytest.mark.parametrize(
    ("data", "expected"),
    [
        (b"\xff\xd8\xff\xe0" + b"\x00" * 64, "image/jpeg"),
        (b"\x89PNG\r\n\x1a\n" + b"\x00" * 64, "image/png"),
        (b"%PDF-1.7\n" + b"\x00" * 64, "application/pdf"),
    ],
)
def test_sniffs_images_and_pdf_by_magic(tmp_path, data, expected):
    assert sniff_mime_type(_write(tmp_path, "Outlook-a1b2c", data)) == expected


def test_sniffs_each_office_zip_by_its_top_folder(tmp_path):
    assert sniff_mime_type(_write(tmp_path, "Outlook-doc", _docx_bytes())) == _DOCX
    assert sniff_mime_type(_write(tmp_path, "Updatepptx", _pptx_bytes())) == _PPTX
    assert sniff_mime_type(_write(tmp_path, "Outlook-xls", _xlsx_bytes())) == _XLSX


def test_a_zip_that_is_not_office_stays_a_zip(tmp_path):
    assert sniff_mime_type(_write(tmp_path, "Outlook-zip", _plain_zip_bytes())) == "application/zip"


def test_sniffs_ole_by_the_streams_it_names(tmp_path):
    assert (
        sniff_mime_type(_write(tmp_path, "a", _ole_bytes("WordDocument"))) == "application/msword"
    )
    assert (
        sniff_mime_type(_write(tmp_path, "b", _ole_bytes("__substg1.0_0037001F")))
        == "application/vnd.ms-outlook"
    )
    assert (
        sniff_mime_type(_write(tmp_path, "c", _ole_bytes("Workbook"))) == "application/vnd.ms-excel"
    )


def test_unknown_bytes_and_missing_files_sniff_to_none(tmp_path):
    assert sniff_mime_type(_write(tmp_path, "blob", b"\x00\x01\x02\x03 random")) is None
    assert sniff_mime_type(str(tmp_path / "missing")) is None


def test_registrar_records_the_sniffed_type_for_an_extensionless_file(tmp_path):
    from src.export.outlook_attachments import register_downloaded_attachments
    from src.store.schema import create_database

    conn = create_database(":memory:")
    try:
        conn.execute(
            "INSERT INTO emails (message_id, date_received, subject) VALUES (?, ?, ?)",
            ("AAMk-sniff", "2026-09-01T10:00:00Z", "synthetic"),
        )
        conn.commit()
        msg_dir = tmp_path / "AAMk-sniff"
        msg_dir.mkdir()
        (msg_dir / "Outlook-img01").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
        (msg_dir / "Outlook-pdf01").write_bytes(b"%PDF-1.7\n" + b"\x00" * 64)
        (msg_dir / "Minutes of the boarddocx").write_bytes(_docx_bytes())
        (msg_dir / "mystery").write_bytes(b"\x00\x01\x02\x03 random")
        (msg_dir / "report.csv").write_bytes(b"a,b\n1,2\n")

        register_downloaded_attachments(conn, tmp_path)

        types = dict(conn.execute("SELECT filename, mime_type FROM attachments").fetchall())
    finally:
        conn.close()
    assert types["Outlook-img01"] == "image/png", "the image backfill selects image/%"
    assert types["Outlook-pdf01"] == "application/pdf"
    assert types["Minutes of the boarddocx"] == _DOCX
    assert types["mystery"] == "application/octet-stream"
    assert types["report.csv"] == "text/csv", "a known extension still decides"


def test_extensionless_docx_labelled_octet_stream_is_extracted(tmp_path):
    path = _write(tmp_path, "Outlook-doc", _docx_bytes())
    result = extract_text_from_file(path, "application/octet-stream")
    assert result["status"] == "extracted"
    assert result["method"] == "python-docx"
    assert "quarterly review" in result["text"]


def test_extensionless_pptx_and_pdf_are_extracted(tmp_path):
    pptx = extract_text_from_file(
        _write(tmp_path, "Status Updatepptx", _pptx_bytes()), "application/octet-stream"
    )
    assert (pptx["status"], pptx["method"]) == ("extracted", "python-pptx")
    pdf = extract_text_from_file(
        _write(tmp_path, "Outlook-pdf", _pdf_bytes()), "application/octet-stream"
    )
    assert (pdf["status"], pdf["method"]) == ("extracted", "pymupdf")


def test_extensionless_xlsx_is_extracted(tmp_path):
    """openpyxl refuses a path whose extension it does not know, so the
    extractor must hand it the bytes rather than the name."""
    result = extract_text_from_file(
        _write(tmp_path, "Outlook-xls", _xlsx_bytes()), "application/octet-stream"
    )
    assert (result["status"], result["method"]) == ("extracted", "openpyxl")


def test_extensionless_image_goes_to_ocr(tmp_path, monkeypatch):
    seen = []
    monkeypatch.setattr(
        ax,
        "_extract_image_ocr",
        lambda p: (
            seen.append(p)
            or {"text": "x", "method": "tesseract", "status": "extracted", "error": None}
        ),
    )
    path = _write(tmp_path, "Outlook-img", b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
    extract_text_from_file(path, "application/octet-stream")
    assert seen == [path]


def test_docx_declared_application_zip_is_extracted_by_its_extension(tmp_path):
    result = extract_text_from_file(
        _write(tmp_path, "brief.docx", _docx_bytes()), "application/zip"
    )
    assert (result["status"], result["method"]) == ("extracted", "python-docx")


def test_extensionless_office_file_declared_application_zip_is_extracted(tmp_path):
    result = extract_text_from_file(
        _write(tmp_path, "Outlook-doc", _docx_bytes()), "application/zip"
    )
    assert (result["status"], result["method"]) == ("extracted", "python-docx")


def test_a_real_archive_is_still_skipped(tmp_path):
    zipped = extract_text_from_file(
        _write(tmp_path, "bundle.zip", _plain_zip_bytes()), "application/zip"
    )
    assert zipped["status"] == "skipped"
    bare = extract_text_from_file(
        _write(tmp_path, "Outlook-zip", _plain_zip_bytes()), "application/zip"
    )
    assert bare["status"] == "skipped"


def test_unknown_bytes_are_still_unsupported(tmp_path):
    result = extract_text_from_file(
        _write(tmp_path, "mystery", b"\x00\x01\x02\x03 random"), "application/octet-stream"
    )
    assert result["status"] == "skipped"
    assert result["error"].startswith("Unsupported type")


def _zip_with_undecodable_name() -> bytes:
    """A zip whose central directory flags its name UTF-8 but holds invalid bytes.

    zipfile decodes a name with bit 0x800 set as UTF-8 while it reads the
    directory, so opening this raises UnicodeDecodeError rather than
    BadZipFile. Any sender can attach one.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("ab", b"payload")
    data = bytearray(buf.getvalue())
    for signature, flag_at, name_at in ((b"PK\x03\x04", 6, 30), (b"PK\x01\x02", 8, 46)):
        start = data.index(signature)
        (flags,) = struct.unpack_from("<H", data, start + flag_at)
        struct.pack_into("<H", data, start + flag_at, flags | 0x800)
        data[start + name_at : start + name_at + 2] = b"\xff\xfe"
    return bytes(data)


def test_a_zip_with_an_undecodable_name_sniffs_to_none(tmp_path):
    path = _write(tmp_path, "Outlook-crafted", _zip_with_undecodable_name())
    with pytest.raises(UnicodeDecodeError):
        zipfile.ZipFile(path)
    assert sniff_mime_type(path) is None


def test_registrar_survives_a_zip_with_an_undecodable_name(tmp_path):
    from src.export.outlook_attachments import register_downloaded_attachments
    from src.store.schema import create_database

    conn = create_database(":memory:")
    try:
        conn.execute(
            "INSERT INTO emails (message_id, date_received, subject) VALUES (?, ?, ?)",
            ("AAMk-crafted", "2026-09-01T10:00:00Z", "synthetic"),
        )
        conn.commit()
        msg_dir = tmp_path / "AAMk-crafted"
        msg_dir.mkdir()
        (msg_dir / "Outlook-crafted").write_bytes(_zip_with_undecodable_name())

        register_downloaded_attachments(conn, tmp_path)

        types = dict(conn.execute("SELECT filename, mime_type FROM attachments").fetchall())
    finally:
        conn.close()
    assert types == {"Outlook-crafted": "application/octet-stream"}


def test_extractor_skips_a_declared_zip_with_an_undecodable_name(tmp_path):
    path = _write(tmp_path, "Outlook-crafted", _zip_with_undecodable_name())
    assert extract_text_from_file(path, "application/zip")["status"] == "skipped"
