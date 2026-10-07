"""The readers that unpack or walk a structure stop at fixed bounds, and say why.

The producer that runs them has 7 GB of memory, shared with the MCP server. A crafted file must
not make a reader inflate without end, read a zip member of any declared size, or walk a
directory chain that loops or never ends.
"""

import io
import struct
import zipfile
import zlib

from src.extract import attachment_extractors as ax
from src.extract.attachment_extractors import _ole_stream_names, extract_text_from_file
from tests.ole_fixtures import (
    RMS_NAMES,
    STREAM,
    mso_file,
    ole_file,
    package_object,
    packed,
)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FOOTER = "Internal use, footer line for every page of the document."


def _write(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    return str(path)


def _xlsx(note: str) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Channel", "Target", note])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- inflation: one budget for every compressed layer of a file ---


def test_an_mso_container_past_the_inflation_budget_is_skipped_with_the_reason(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(ax, "INFLATE_MAX_BYTES", 1024)
    data = packed(ole_file([("_1", STREAM, b"\x00" * 8192)]))

    out = extract_text_from_file(_write(tmp_path, "oledata.mso", data), "")

    assert (out["status"], out["method"]) == ("skipped", "mso")
    assert "inflates past 1,024 bytes" in out["error"]
    assert "file kept" in out["error"]


def test_the_budget_covers_all_the_objects_of_a_container_together(tmp_path, monkeypatch):
    """Each object alone fits; the second one would take the file past the budget."""
    first = package_object(_xlsx("The first object is read in full, it fits the budget"))
    second = package_object(_xlsx("The second object would cross the budget and is not read"))
    data = mso_file([first, second])
    outer = zlib.decompress(data[4:])
    monkeypatch.setattr(ax, "INFLATE_MAX_BYTES", len(outer) + len(first) + 100)

    out = extract_text_from_file(_write(tmp_path, "oledata.mso", data), "")

    assert out["status"] == "extracted"
    assert "The first object is read in full" in out["text"]
    assert "cross the budget" not in out["text"]
    assert "embedded object 2: inflates past" in out["error"]


def test_inflation_asks_zlib_for_no_more_than_the_budget(monkeypatch):
    """decompressobj with max_length, never zlib.decompress: the output cannot outgrow it."""
    budget = ax._InflateBudget(4096)
    bomb = zlib.compress(b"\x00" * (50 << 20))

    try:
        budget.inflate(bomb)
        raised = False
    except ax._InflatesTooFar:
        raised = True

    assert raised


# --- the XML fallbacks read a member only within its size and ratio bounds ---


def _docx_with_null_footer(body: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(body)
    doc.sections[0].footer.paragraphs[0].text = FOOTER
    buf = io.BytesIO()
    doc.save(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                blob = blob.replace(b'Target="footer1.xml"', b'Target="NULL"')
            dst.writestr(info.filename, blob)
    return out.getvalue()


def test_the_docx_fallback_refuses_a_member_past_the_size_bound(tmp_path, monkeypatch):
    body = "A body paragraph long enough to be larger than the bound set here. " * 40
    path = _write(tmp_path, "Plan.docx", _docx_with_null_footer(body))
    with zipfile.ZipFile(path) as zf:
        document = zf.getinfo("word/document.xml").file_size
        footer = zf.getinfo("word/footer1.xml").file_size
    monkeypatch.setattr(ax, "XML_PART_MAX_BYTES", (document + footer) // 2)
    assert footer < ax.XML_PART_MAX_BYTES < document

    out = extract_text_from_file(path, DOCX)

    assert out["method"] == "docx-xml"
    assert FOOTER in out["text"] and "A body paragraph" not in out["text"]


def test_the_docx_fallback_refuses_a_member_compressed_like_a_bomb(tmp_path):
    body = "x" * 400_000  # deflates far past the archive ratio bound
    path = _write(tmp_path, "Plan.docx", _docx_with_null_footer(body))
    with zipfile.ZipFile(path) as zf:
        info = zf.getinfo("word/document.xml")
    assert info.file_size > ax.ZIP_MAX_RATIO * info.compress_size

    out = extract_text_from_file(path, DOCX)

    assert out["method"] == "docx-xml"
    assert FOOTER in out["text"] and "xxxxxxxx" not in out["text"]


def test_a_member_read_stops_at_the_bound_whatever_its_header_says(tmp_path, monkeypatch):
    """A header can understate a member's size; the read itself is capped too."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/document.xml", "<w:p><w:t>" + "y" * 5000 + "</w:t></w:p>")
    monkeypatch.setattr(ax, "XML_PART_MAX_BYTES", 1000)
    with zipfile.ZipFile(buf) as zf:
        zf.getinfo("word/document.xml").file_size = 10  # what a lying header would say
        assert ax._read_member(zf, "word/document.xml") is None


# --- the OLE directory walk ---


def test_the_directory_walk_stops_at_the_entry_bound(tmp_path, monkeypatch):
    monkeypatch.setattr(ax, "_OLE_MAX_DIRECTORY_ENTRIES", 4)
    path = _write(tmp_path, "book.xlsx", ole_file(RMS_NAMES))

    names = _ole_stream_names(path)

    assert "\x06DataSpaces" in names  # the first sector, four entries
    assert "EncryptedPackage" not in names


def test_a_difat_chain_that_loops_ends(tmp_path):
    """The header claims more FAT sectors than it lists, and the DIFAT chain points at itself."""
    raw = bytearray(ole_file(RMS_NAMES, pad_sectors=1))
    pad = 1  # the FAT is sector 0, the padding sector 1
    struct.pack_into("<II", raw, 44, 300, 2)  # 300 FAT sectors; directory still at sector 2
    struct.pack_into("<II", raw, 68, pad, 5)  # DIFAT starts at the padding sector, 5 of them
    block = [0xFFFFFFFF] * 127 + [pad]  # its next-DIFAT pointer is itself
    struct.pack_into("<128I", raw, 512 * (pad + 1), *block)

    names = _ole_stream_names(_write(tmp_path, "book.xlsx", bytes(raw)))

    assert "EncryptedPackage" in names
