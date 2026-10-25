"""Figures read from a recurring report's stored cell dump, with no model call.

Numbers were 12 of the owner's 93 questions in six weeks, the costliest at a median of 12
brain calls each, and four of his twelve complaints: the figures sit in recurring workbooks,
stored as a cell dump that agents searched with LIKE. figures_parse reads that dump into rows of
(sheet, row label, column header, period, value, unit, cell). Every workbook here is made with
openpyxl and read by the real Phase 1 reader, in the layouts recurring reports use: dates as
columns, months as columns, months down the rows under a year row, dates down the rows, and a
period named only in the filename or by the email's date.
"""

from datetime import date, datetime

import openpyxl
import pytest

from src.extract.attachment_extractors import extract_text_from_file
from src.extract.figures_parse import (
    filename_period,
    parse_period,
    parse_workbook,
    series_stem,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAILED = "2026-10-12T09:30:00Z"


def _dump(tmp_path, sheets: dict[str, list[list]]) -> str:
    """The text Phase 1 stores for a workbook holding these sheets."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    path = tmp_path / "book.xlsx"
    wb.save(path)
    result = extract_text_from_file(str(path), XLSX)
    assert result["status"] == "extracted"
    return result["text"]


def _by(figures, **match) -> list:
    return [f for f in figures if all(getattr(f, k) == v for k, v in match.items())]


def _one(figures, **match):
    found = _by(figures, **match)
    assert len(found) == 1, (match, found)
    return found[0]


# ------------------------------------------------------------------ series


@pytest.mark.parametrize(
    "name",
    [
        "Monthly KPI Report 31.08.2026.xlsx",
        "Monthly KPI Report 30.09.2026 v2.xlsx",
        "Copy of Monthly KPI Report 31-07-2026 (1).xlsx",
        "20260831_Monthly KPI Report_final.xlsx",
        "Monthly_KPI_Report_Aug-26.xlsx",
        "Monthly KPI Report Q3 2026 - Copy.xlsx",
        "Monthly KPI Report 202608.xlsb",
        "MONTHLY KPI REPORT 2026 08 31.xls",
        "Monthly KPI Report 2026 08.xlsx",
    ],
)
def test_editions_share_a_stem_whatever_their_date_version_or_copy_marker(name):
    assert series_stem(name) == "monthly kpi report"


def test_greek_month_names_and_accents_fold_into_one_stem():
    stems = {
        series_stem("Έκθεση Πωλήσεων Αύγουστος 2026.xlsx"),
        series_stem("ΕΚΘΕΣΗ πωλήσεων Σεπτεμβρίου 2026.xlsx"),
        series_stem("Έκθεση Πωλήσεων - 28 Αυγ 2026.xlsx"),
    }
    assert stems == {"εκθεση πωλησεων"}


def test_a_code_that_is_no_date_keeps_two_files_apart():
    """Six digits that read as no month are a code: one file per branch is not a series."""
    assert series_stem("717208-DEP-20260616.xlsx") != series_stem("717234-DEP-20260616.xlsx")


def test_a_name_that_is_only_a_date_has_an_empty_stem():
    assert series_stem("02-04-2026.xlsx") == ""


@pytest.mark.parametrize(
    ("name", "label"),
    [
        ("Monthly KPI Report 31.08.2026.xlsx", "2026-08-31"),
        ("Daily Summary 02042026.xlsx", "2026-04-02"),
        ("20260831_Daily report.xlsx", "2026-08-31"),
        ("Report Aug-26.xlsx", "2026-08"),
        ("Έκθεση Αύγουστος 2026.xlsx", "2026-08"),
        ("Report 2026-09.xlsx", "2026-09"),
        ("Report Q3 2026.xlsx", "2026-Q3"),
        ("Transactions 01Jan2024-31Oct2025.xls", "2025-10-31"),
        ("Budget 2027.xlsx", "2027"),
    ],
)
def test_the_filename_names_a_period(name, label):
    assert filename_period(name).label == label


def test_a_filename_without_a_date_names_none():
    assert filename_period("Benefit accounts daily.xlsx") is None


@pytest.mark.parametrize(
    ("text", "label", "end"),
    [
        ("2026-08-31 00:00:00", "2026-08-31", "2026-08-31"),
        ("31/08/2026", "2026-08-31", "2026-08-31"),
        ("Αύγουστος 2026", "2026-08", "2026-08-31"),
        ("Aug-26", "2026-08", "2026-08-31"),
        ("Φεβ 2024", "2024-02", "2024-02-29"),
        ("Q3 2026", "2026-Q3", "2026-09-30"),
        ("H1 2026", "2026-H1", "2026-06-30"),
        ("2025", "2025", "2025-12-31"),
    ],
)
def test_a_period_reads_from_dates_months_quarters_and_years(text, label, end):
    found = parse_period(text)
    assert found is not None
    assert (found.period.label, found.period.end) == (label, end)


def test_a_month_without_a_year_takes_the_year_it_is_given():
    found = parse_period("Μάρτιος", year=2025)
    assert found is not None and found.period.label == "2025-03"


def test_words_beside_a_period_are_its_rest():
    found = parse_period("YTD Aug 2026")
    assert found is not None
    assert (found.period.label, found.rest) == ("2026-08", "ytd")


# ------------------------------------------------------------------ layouts


def test_dates_as_columns(tmp_path):
    month_ends = [datetime(2026, 7, 31), datetime(2026, 8, 31), datetime(2026, 9, 30)]
    text = _dump(
        tmp_path,
        {
            "Payments": [
                ["Metric", *month_ends],
                ["Card payments", 100, 110, 120],
                ["Wallet payments", 7, 8, 9],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Payments.xlsx", email_date=MAILED)

    assert len(parsed.figures) == 6
    september = _one(parsed.figures, label="Card payments", period="2026-09")
    assert (september.value, september.period_source, september.header_key) == (
        120,
        "header",
        "",
    )
    assert september.period_end == "2026-09-30"


def test_months_as_columns_take_the_year_of_the_filename(tmp_path):
    text = _dump(
        tmp_path,
        {"KPIs": [["KPI", "Ιανουάριος", "Φεβρουάριος", "Μάρτιος"], ["Νέοι πελάτες", 10, 12, 15]]},
    )

    parsed = parse_workbook(text, filename="KPI report 2025.xlsx", email_date=MAILED)

    assert [(f.period, f.value) for f in parsed.figures] == [
        ("2025-01", 10),
        ("2025-02", 12),
        ("2025-03", 15),
    ]


def test_months_down_the_rows_take_the_year_row_above_them(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Index": [
                ["Dates", "Branch", "Digital", "% Index"],
                [2025],
                ["Νοέμβριος", 80, 20, 0.2],
                ["Δεκέμβριος", 70, 30, 0.3],
                [2026],
                ["Ιανουάριος", 60, 40, 0.4],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Index 31.01.2026.xlsx", email_date=MAILED)

    index = _by(parsed.figures, header="% Index")
    assert [(f.period, f.value, f.period_source) for f in index] == [
        ("2025-11", 0.2, "row"),
        ("2025-12", 0.3, "row"),
        ("2026-01", 0.4, "row"),
    ]
    assert {f.unit for f in index} == {"%"}
    assert {f.label_key for f in index} == {""}


def test_dates_down_the_rows(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Daily": [
                ["Date", "New accounts", "Closed accounts"],
                [date(2026, 10, 7), 675, 742],
                [date(2026, 10, 8), 932, 1286],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Accounts daily.xlsx", email_date=MAILED)

    new = _by(parsed.figures, header="New accounts")
    assert [(f.period, f.value) for f in new] == [("2026-10-07", 675), ("2026-10-08", 932)]
    assert {f.period_source for f in new} == {"row"}


def test_day_and_month_without_a_year_take_the_editions_year(tmp_path):
    text = _dump(
        tmp_path,
        {"Snapshots": [["Date", "Opened"], ["30/09", 5], ["01/10", 6], ["31/12", 4]]},
    )

    parsed = parse_workbook(text, filename="Snapshots.xlsx", email_date=MAILED)

    # Mailed in October 2026: a 31 December ahead of it is last year's.
    assert [f.period for f in parsed.figures] == ["2026-09-30", "2026-10-01", "2025-12-31"]


def test_a_period_only_in_the_filename(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Lines": [
                ["Line", "Entered", "Answered", "% Service level"],
                ["Retail line", 1944, 1909, 0.83],
                ["Business line", 152, 146, 0.88],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Call centre 30.09.2026.xlsx", email_date=MAILED)

    assert len(parsed.figures) == 6
    level = _one(parsed.figures, label="Retail line", header="% Service level")
    assert (level.value, level.unit, level.period, level.period_source) == (
        0.83,
        "%",
        "2026-09-30",
        "filename",
    )


def test_the_emails_date_when_nothing_else_names_one(tmp_path):
    text = _dump(tmp_path, {"Lines": [["Line", "Entered"], ["Retail line", 1944]]})

    parsed = parse_workbook(text, filename="Call centre.xlsx", email_date=MAILED)

    assert [(f.period, f.period_source) for f in parsed.figures] == [("2026-10-12", "email")]


def test_a_date_in_the_sheet_name_or_above_the_table_beats_the_filename(tmp_path):
    text = _dump(
        tmp_path,
        {
            "14.10.2026": [["Line", "Entered"], ["Retail line", 1944]],
            "Other": [[datetime(2026, 10, 13)], ["Line", "Entered"], ["Retail line", 1800]],
        },
    )

    parsed = parse_workbook(text, filename="Call centre 15.10.2026.xlsx", email_date=MAILED)

    assert [(f.sheet, f.period, f.period_source) for f in parsed.figures] == [
        ("14.10.2026", "2026-10-14", "sheet"),
        ("Other", "2026-10-13", "sheet"),
    ]


def test_a_year_by_month_matrix(tmp_path):
    text = _dump(
        tmp_path,
        {"Volumes": [["Year", "Jan", "Feb"], [2025, 10, 11], [2026, 20, 21]]},
    )

    parsed = parse_workbook(text, filename="Volumes.xlsx", email_date=MAILED)

    assert [(f.period, f.value) for f in parsed.figures] == [
        ("2025-01", 10),
        ("2025-02", 11),
        ("2026-01", 20),
        ("2026-02", 21),
    ]


def test_a_value_row_with_no_label_under_dates_is_named_by_the_sheets_title(tmp_path):
    days = [datetime(2026, 9, 29), datetime(2026, 9, 30)]
    text = _dump(
        tmp_path,
        {
            "Count": [
                ["Total number of e-payments"],
                days,
                [5616502, 6598449],
                [datetime(2025, 9, 29), datetime(2025, 9, 30)],
                [4902465, 5987988],
            ]
        },
    )

    parsed = parse_workbook(text, filename="e-payments.xlsx", email_date=MAILED)

    assert [(f.label, f.period, f.value) for f in parsed.figures] == [
        ("Total number of e-payments", "2026-09-29", 5616502),
        ("Total number of e-payments", "2026-09-30", 6598449),
        ("Total number of e-payments", "2025-09-29", 4902465),
        ("Total number of e-payments", "2025-09-30", 5987988),
    ]


def test_placeholders_are_neither_values_nor_labels(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Lines": [
                ["Line", "Entered", "Answered"],
                ["Retail line", "-", 1909],
                ["Cards", "n/a", 5],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Lines 30.09.2026.xlsx", email_date=MAILED)

    assert [(f.label, f.header, f.value) for f in parsed.figures] == [
        ("Retail line", "Answered", 1909),
        ("Cards", "Answered", 5),
    ]


def test_a_row_with_an_empty_cell_is_left_out_rather_than_guessed(tmp_path):
    """The dump drops an empty cell, so a short row cannot be placed under its headers."""
    text = _dump(
        tmp_path,
        {
            "Lines": [
                ["Line", "Entered", "Answered", "Lost"],
                ["Retail", 10, None, 2],
                ["Cards", 5, 4, 1],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Lines 30.09.2026.xlsx", email_date=MAILED)

    assert {f.label for f in parsed.figures} == {"Cards"}
    assert parsed.skipped["unaligned_rows"] == 1


def test_a_group_label_shown_once_carries_down_to_the_rows_under_it(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Sales": [
                ["Family", "Product", "Units"],
                ["Cards", "Credit", 100],
                [None, "Debit", 50],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Sales 30.09.2026.xlsx", email_date=MAILED)

    assert [(f.label, f.value) for f in parsed.figures] == [
        ("Cards / Credit", 100),
        ("Cards / Debit", 50),
    ]


def test_a_section_row_names_the_rows_under_it(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Sales": [
                ["Product", "Units"],
                ["Individuals"],
                ["Debit cards", 100],
                ["Business"],
                ["Debit cards", 7],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Sales 30.09.2026.xlsx", email_date=MAILED)

    assert [(f.label, f.value) for f in parsed.figures] == [
        ("Individuals / Debit cards", 100),
        ("Business / Debit cards", 7),
    ]


def test_a_header_split_by_line_breaks_is_joined_again(tmp_path):
    text = _dump(
        tmp_path,
        {"Sales": [["Month", "Digital\nsales", "Total\nsales"], [datetime(2026, 8, 31), 5, 20]]},
    )

    parsed = parse_workbook(text, filename="Sales.xlsx", email_date=MAILED)

    assert [(f.header, f.value) for f in parsed.figures] == [
        ("Digital sales", 5),
        ("Total sales", 20),
    ]


def test_a_repeated_header_is_told_apart_by_its_place(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Sales": [
                ["Month", "Total", "Index", "Total", "Index"],
                [datetime(2026, 8, 31), 5, 0.1, 9, 0.2],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Sales.xlsx", email_date=MAILED)

    assert len({(f.header_key, f.period) for f in parsed.figures}) == 4


def test_a_record_list_is_not_a_table_of_figures(tmp_path):
    rows = [["Customer", "Balance"]] + [[f"Customer {i}", i] for i in range(1, 600)]
    text = _dump(tmp_path, {"Data": rows, "Summary": [["Line", "Total"], ["All", 599]]})

    parsed = parse_workbook(text, filename="Book 30.09.2026.xlsx", email_date=MAILED)

    assert {f.sheet for f in parsed.figures} == {"Summary"}
    assert parsed.skipped["record_lists"] == 1


def test_every_figure_points_at_its_cell_in_the_stored_text(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Notes": [["A note on the report"]],
            "Lines": [
                ["Line", "Entered", "Answered"],
                ["Retail line", 1944, 1909],
                ["Cards", 5, 4],
            ],
        },
    )

    parsed = parse_workbook(text, filename="Lines 30.09.2026.xlsx", email_date=MAILED)

    answered = _one(parsed.figures, label="Cards", header="Answered")
    assert answered.cell == "R3C3"
    assert text[answered.offset :].startswith("4")
    for figure in parsed.figures:
        stored = text[figure.offset :].split(" | ", 1)[0].split("\n", 1)[0]
        assert float(stored) == figure.value


def test_a_row_with_words_after_its_figures_is_left_out(tmp_path):
    """A code read as a figure splits the labels: which name goes with which figure is unknown."""
    text = _dump(
        tmp_path,
        {
            "Branches": [
                ["Region", "Code", "Branch", "Inflows", "Outflows"],
                ["North", 46, "Harbour street", 25000, 10],
                ["North", "Total", "All branches", 90000, 30],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Branches 30.09.2026.xlsx", email_date=MAILED)

    assert {f.label for f in parsed.figures} == {"North / Total / All branches"}
    assert parsed.skipped["unaligned_rows"] == 1


def test_rows_led_by_many_labels_are_records_not_figures(tmp_path):
    rows = [["Order", "Branch", "Kind", "Status", "Amount"]]
    rows += [[f"Order {i}", f"Branch {i % 7}", "Supply", "Confirmed", 1000 + i] for i in range(20)]
    text = _dump(tmp_path, {"Orders": rows})

    parsed = parse_workbook(text, filename="Orders 30.09.2026.xlsx", email_date=MAILED)

    assert parsed.figures == []
    assert parsed.skipped["record_lists"] == 1


def test_excel_errors_are_placeholders(tmp_path):
    text = _dump(tmp_path, {"Lines": [["Line", "Entered", "Rate"], ["Retail", 10, "#DIV/0!"]]})

    parsed = parse_workbook(text, filename="Lines 30.09.2026.xlsx", email_date=MAILED)

    assert [(f.header, f.value) for f in parsed.figures] == [("Entered", 10)]


def test_a_to_date_header_is_as_of_the_report_not_the_year(tmp_path):
    text = _dump(
        tmp_path,
        {"Funds": [["Product", "MtD 2026", "YtD 2026"], ["Mutual funds", 120, 900]]},
    )

    parsed = parse_workbook(text, filename="Funds 2026 10 14.xlsx", email_date=MAILED)

    assert [(f.header_key, f.period, f.period_source) for f in parsed.figures] == [
        ("mtd", "2026-10-14", "filename"),
        ("ytd", "2026-10-14", "filename"),
    ]


def test_rows_named_by_their_period_are_labelled_by_the_sheets_title(tmp_path):
    """Dates down the rows and products across: each row's label is its period, so the label
    the figure is shown and found by is the sheet's title, its first row of one cell. Its key
    stays empty, so a title reworded between editions keeps one metric."""
    text = _dump(
        tmp_path,
        {
            "KYC": [
                ["Digital migration index KYC"],
                ["Retail reporting"],
                ["Month", "Upload", "Total"],
                [datetime(2026, 6, 1), 10, 40],
                [datetime(2026, 7, 1), 12, 41],
                [datetime(2026, 8, 1), 14, 42],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Index 31.08.2026.xlsx", email_date=MAILED)

    upload = _by(parsed.figures, header="Upload")
    assert {f.label for f in upload} == {"Digital migration index KYC"}
    assert {f.label_key for f in upload} == {""}
    assert [f.period for f in upload] == ["2026-06", "2026-07", "2026-08"]


def test_without_a_title_the_sheet_name_labels_them(tmp_path):
    text = _dump(
        tmp_path,
        {"Consumer loans": [["Month", "Digital", "Total"], [datetime(2026, 8, 1), 5, 7]]},
    )

    parsed = parse_workbook(text, filename="Index 31.08.2026.xlsx", email_date=MAILED)

    assert {f.label for f in parsed.figures} == {"Consumer loans"}


def test_a_section_names_the_period_rows_under_it(tmp_path):
    text = _dump(
        tmp_path,
        {
            "Sales": [
                ["Sales index"],
                ["Month", "Cards"],
                ["Individuals"],
                [datetime(2026, 8, 1), 5],
                ["Business"],
                [datetime(2026, 8, 1), 2],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Index 31.08.2026.xlsx", email_date=MAILED)

    assert [(f.label, f.value) for f in parsed.figures] == [("Individuals", 5), ("Business", 2)]


def test_a_row_with_no_label_takes_the_label_above_it_in_its_column(tmp_path):
    """A product label merged down over two segment blocks is stored once: the row under the
    next segment has its figures and no label of its own."""
    text = _dump(
        tmp_path,
        {
            "Sales": [
                ["Product", "Units", "Value"],
                ["Individuals"],
                ["Cards", 10, 100],
                ["Business"],
                [None, 7, 70],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Sales 30.09.2026.xlsx", email_date=MAILED)

    assert [(f.label, f.header, f.value) for f in parsed.figures] == [
        ("Individuals / Cards", "Units", 10),
        ("Individuals / Cards", "Value", 100),
        ("Business / Cards", "Units", 7),
        ("Business / Cards", "Value", 70),
    ]


def test_text_that_is_no_workbook_dump_gives_nothing():
    parsed = parse_workbook("Just a paragraph.", filename="x.xlsx", email_date=MAILED)
    assert parsed.figures == []


def _merged_report(tmp_path) -> str:
    """A paginated report exported to Excel: a title merged over
    the date column, three levels of column groups in merged cells over "€ | #" pairs, a group
    with no level under it, and one value merged over two columns in its header and in every
    data row alike."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily turnover"
    ws["A1"] = "DAILY CARD MOVEMENTS"
    ws.merge_cells("A1:F1")
    ws["H1"] = "Complete until"
    ws["J1"] = datetime(2026, 10, 2)
    for cell, text, span in [
        ("B2", "DEBIT", "B2:H2"),
        ("I2", "CREDIT", "I2:L2"),
        ("B3", "Purchases", "B3:F3"),
        ("G3", "Refunds", "G3:H3"),
        ("I3", "Purchases", "I3:J3"),
        ("K3", "Instalments", "K3:L3"),
        ("B4", "Domestic", "B4:C4"),
        ("D4", "Abroad", "D4:F4"),
        ("G4", "Domestic", "G4:H4"),
        ("I4", "Domestic", "I4:J4"),
    ]:
        ws[cell] = text
        ws.merge_cells(span)
    ws.append(["Date", "€", "#", "€", None, "#", "€", "#", "€", "#", "€", "#"])
    ws.merge_cells("D5:E5")
    for r, day in ((6, 1), (7, 2)):
        ws.append(
            [datetime(2026, 10, day), 100 * day, day, 50 * day, None, 5 * day, 9, 1, 70, 7, 30, 3]
        )
        ws.merge_cells(f"D{r}:E{r}")
    path = tmp_path / "Daily card turnover.xlsx"
    wb.save(path)
    result = extract_text_from_file(str(path), XLSX)
    assert result["status"] == "extracted"
    return result["text"]


def test_merged_column_groups_name_each_value_column(tmp_path):
    """Each value column is named by every merged group above it, down to its own "€" or
    "#": a header of "€" alone leaves every figure of such a report unsearchable."""
    text = _merged_report(tmp_path)

    parsed = parse_workbook(text, filename="Daily card turnover.xlsx", email_date=MAILED)

    day1 = {f.header: f.value for f in parsed.figures if f.period == "2026-10-01"}
    assert day1 == {
        "DEBIT / Purchases / Domestic / €": 100,
        "DEBIT / Purchases / Domestic / #": 1,
        "DEBIT / Purchases / Abroad / €": 50,
        "DEBIT / Purchases / Abroad / #": 5,
        "DEBIT / Refunds / Domestic / €": 9,
        "DEBIT / Refunds / Domestic / #": 1,
        "CREDIT / Purchases / Domestic / €": 70,
        "CREDIT / Purchases / Domestic / #": 7,
        "CREDIT / Instalments / €": 30,
        "CREDIT / Instalments / #": 3,
    }
    assert parsed.skipped["unaligned_rows"] == 0


def test_a_total_row_under_dates_down_the_rows_is_no_day_of_its_own(tmp_path):
    """The sum row under a run of days has no date: dated by the email it read as the newest
    day, so the month's total answered for "the latest day"."""
    text = _dump(
        tmp_path,
        {
            "Daily": [
                ["Date", "Purchases", "Refunds"],
                [datetime(2026, 10, 1), 100, 1],
                [datetime(2026, 10, 2), 120, 2],
                [None, 220, 3],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Daily cards.xlsx", email_date=MAILED)

    assert sorted({f.period for f in parsed.figures}) == ["2026-10-01", "2026-10-02"]
    assert parsed.skipped["totals"] == 1


def test_tables_side_by_side_under_one_header_row_keep_their_months(tmp_path):
    """Two year blocks side by side, each led by the same corner header, over rows led by a
    group label named once: the dump drops the empty cell over the group column and the
    right block's empty label cells, so the first row is as wide as the header while its
    figures sit one column right of their months."""
    this_year = [datetime(2026, m, d) for m, d in ((1, 31), (2, 28), (3, 31))]
    last_year = [datetime(2025, m, d) for m, d in ((10, 31), (11, 30), (12, 31))]
    corner = "Based on application date"
    text = _dump(
        tmp_path,
        {
            "Dashboard": [
                [None, corner, *this_year, "YTD", corner, *last_year, "YTD"],
                ["Items", "Total applications", 10, 11, 12, 33, None, 20, 21, 22, 63],
                [None, "Pending", 1, 2, 3, 6, None, 4, 5, 6, 15],
                [None, "Booked", 7, 8, 9, 24, None, 1, 1, 1, 3],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Applications_2026-03-21.xlsx", email_date=MAILED)

    def months(label):
        return {f.period: f.value for f in parsed.figures if f.label == label and not f.header}

    assert months("Items / Total applications") == {
        "2026-01": 10,
        "2026-02": 11,
        "2026-03": 12,
        "2025-10": 20,
        "2025-11": 21,
        "2025-12": 22,
    }
    assert months("Items / Pending") == {
        "2026-01": 1,
        "2026-02": 2,
        "2026-03": 3,
        "2025-10": 4,
        "2025-11": 5,
        "2025-12": 6,
    }
    totals = {f.value for f in parsed.figures if f.label == "Items / Booked" and f.header}
    assert totals == {24, 3}
    assert all(corner.lower() not in f.header.lower() for f in parsed.figures)
    assert parsed.skipped["unaligned_rows"] == 0


def test_a_group_row_is_not_joined_to_a_header_short_of_empty_cells(tmp_path):
    """A row of column groups over a header with two empty cells: joined at their boundary
    they are exactly as wide as the data, which read every figure two columns left of its
    header. The header is not split by a line break, so the table has no header to read."""
    quarters = ["Q2 2026", "Q1 2026", "Q4 2025"]
    text = _dump(
        tmp_path,
        {
            "Evolution": [
                ["Market share evolution"],
                [None, "Market", None, None, None, "Ours", None, None, "Our share"],
                ["Cards ('000s)", None, *quarters, None, *quarters[:2], "% Q2 2026"],
                ["Debit", 1.5, 100, 90, 80, 2.5, 30, 28, 0.3],
                ["Credit", 0.5, 50, 45, 40, 1.5, 10, 9, 0.2],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Cards Q2 2026.xlsx", email_date=MAILED)

    assert parsed.figures == []
    assert parsed.skipped["no_header"] == 1


def test_a_group_row_over_a_header_as_wide_as_the_data_names_the_columns(tmp_path):
    """The same groups over a header with no empty cell: the header is read as it stands."""
    text = _dump(
        tmp_path,
        {
            "Evolution": [
                [None, "Market", None, "Ours", None],
                ["Cards ('000s)", "Q2 2026", "Q1 2026", "% D qoq", "Q4 2025"],
                ["Debit", 100, 90, 0.1, 80],
            ]
        },
    )

    parsed = parse_workbook(text, filename="Cards Q2 2026.xlsx", email_date=MAILED)

    assert sorted((f.header, f.period, f.value) for f in parsed.figures) == [
        ("", "2025-Q4", 80),
        ("", "2026-Q1", 90),
        ("", "2026-Q2", 100),
        ("% D qoq", "2026-Q2", 0.1),
    ]
