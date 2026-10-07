"""Office files that python-docx or python-pptx refuse, but whose text is in the zip.

python-docx loads every relationship target, and Word writes some footer relationships with
Target="NULL", so Document() raised KeyError ("no item named 'word/NULL'") on 9 files whose body
was intact. python-pptx reads every part, so one media member with a bad CRC failed 4 decks
whose slides were fine. Both now fall back to reading the text runs straight from the XML. A
file named .pptx that is not a zip at all (one was an HTML page) goes by its bytes instead.
"""

import io
import struct
import zipfile

import pytest

from src.extract import attachment_extractors as ax
from src.extract.attachment_extractors import extract_text_from_file

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
BODY = "The regional plan for the coming quarter, written out in plenty of words."
FOOTER = "Internal use, footer line for every page of the document."
SLIDE = "Network review: branches, channels and the targets for next year."
NOTES = "Speaker notes that explain the second chart in some detail."


def _rewrite(data: bytes, change) -> bytes:
    """Copy a zip member by member, letting `change(name, bytes)` alter one."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, change(info.filename, src.read(info.filename)))
    return out.getvalue()


def _docx_with_null_footer() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(BODY)
    doc.sections[0].footer.paragraphs[0].text = FOOTER
    buf = io.BytesIO()
    doc.save(buf)

    def null_target(name, blob):
        if name == "word/_rels/document.xml.rels":
            return blob.replace(b'Target="footer1.xml"', b'Target="NULL"')
        return blob

    return _rewrite(buf.getvalue(), null_target)


def _png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, format="PNG")
    return buf.getvalue()


def _pptx(notes: bool = True) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = SLIDE
    slide.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(3))
    if notes:
        slide.notes_slide.notes_text_frame.text = NOTES
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _break_crc(data: bytes, member_prefix: str) -> bytes:
    """Write a wrong CRC-32 for one member, in its local header and the central directory."""
    raw = bytearray(data)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        info = next(i for i in zf.infolist() if i.filename.startswith(member_prefix))
    bad = info.CRC ^ 0xFFFFFFFF
    struct.pack_into("<I", raw, info.header_offset + 14, bad)
    (entry,) = struct.unpack_from("<I", raw, raw.rindex(b"PK\x05\x06") + 16)
    while raw[entry : entry + 4] == b"PK\x01\x02":
        name_len, extra_len, comment_len = struct.unpack_from("<HHH", raw, entry + 28)
        if raw[entry + 46 : entry + 46 + name_len].decode() == info.filename:
            struct.pack_into("<I", raw, entry + 16, bad)
            return bytes(raw)
        entry += 46 + name_len + extra_len + comment_len
    raise AssertionError(f"no central directory entry for {info.filename}")


def _write(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    return str(path)


def test_a_docx_with_a_null_relationship_target_is_read_from_its_xml(tmp_path):
    path = _write(tmp_path, "Plan.docx", _docx_with_null_footer())
    with pytest.raises(KeyError):
        from docx import Document

        Document(path)

    out = extract_text_from_file(path, DOCX)

    assert (out["status"], out["method"]) == ("extracted", "docx-xml")
    assert BODY in out["text"] and FOOTER in out["text"]


def test_a_pptx_with_a_bad_crc_media_member_is_read_from_its_xml(tmp_path):
    path = _write(tmp_path, "Review.pptx", _break_crc(_pptx(), "ppt/media/"))
    with pytest.raises(zipfile.BadZipFile):
        from pptx import Presentation

        Presentation(path)

    out = extract_text_from_file(path, PPTX)

    assert (out["status"], out["method"]) == ("extracted", "pptx-xml")
    assert "--- Slide 1 ---" in out["text"]
    assert SLIDE in out["text"]
    assert f"[Notes] {NOTES}" in out["text"]


def test_a_readable_deck_still_goes_through_python_pptx(tmp_path):
    out = extract_text_from_file(_write(tmp_path, "Review.pptx", _pptx()), PPTX)

    assert (out["status"], out["method"]) == ("extracted", "python-pptx")


def test_a_truncated_deck_is_still_a_failure(tmp_path):
    data = _pptx()
    out = extract_text_from_file(_write(tmp_path, "Cut.pptx", data[: len(data) // 2]), PPTX)

    assert out["status"] == "failed"


def test_an_html_page_named_pptx_and_labelled_html_is_read_as_html(tmp_path):
    """The real case: the name said .pptx, the declared type said text/html, and the name won."""
    page = f"<html><head><title>Review</title></head><body><p>{SLIDE}</p></body></html>"
    out = extract_text_from_file(_write(tmp_path, "Review.pptx", page.encode()), "text/html")

    assert (out["status"], out["method"]) == ("extracted", "direct_read")
    assert SLIDE in out["text"] and "<p>" not in out["text"]


def test_a_docx_with_one_corrupt_member_keeps_the_text_of_the_others(tmp_path):
    """A bad member that is not the body (here the footer) costs only its own text."""
    data = _break_crc(_docx_with_null_footer(), "word/footer")
    out = extract_text_from_file(_write(tmp_path, "Plan.docx", data), DOCX)

    assert (out["status"], out["method"]) == ("extracted", "docx-xml")
    assert BODY in out["text"] and FOOTER not in out["text"]


def test_a_file_named_pptx_that_is_not_a_zip_is_not_failed_as_one(tmp_path):
    """The bytes decide (tests/test_sniffed_formats.py reads an HTML page named .pptx)."""
    out = extract_text_from_file(_write(tmp_path, "Odd.pptx", b"\x00\x01\x02 not a zip"), PPTX)

    assert out["status"] == "skipped"
    assert "PackageNotFoundError" not in (out["error"] or "")


def test_a_legacy_document_named_docx_goes_to_the_doc_reader(tmp_path, monkeypatch):
    from tests.ole_fixtures import STREAM, ole_file

    seen = []
    monkeypatch.setattr(
        ax,
        "_extract_doc",
        lambda p: (
            seen.append(p)
            or {"text": BODY, "method": "antiword", "status": "extracted", "error": None}
        ),
    )
    path = _write(tmp_path, "Old.docx", ole_file([("WordDocument", STREAM), ("1Table", STREAM)]))

    out = extract_text_from_file(path, DOCX)

    assert seen == [path]
    assert out["method"] == "antiword"
